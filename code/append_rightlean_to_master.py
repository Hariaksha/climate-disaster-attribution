"""
Append the cleaned right-leaning candidates to the Master sheet of
attribution.xlsx.

Input : data/rightlean_articles_clean.csv   (733 rows)
Edits : attribution.xlsx -- Master (append only), plus formula-range
        extensions on All Disasters, Active and Trash.

Design notes
------------
Events are keyed on (Name, Begin Date), not Name alone. Three event names
occur twice in All Disasters -- "Midwest Severe Weather (June 2017)",
"Colorado Hail Storm (June 2018)" and "North Central Hail Storms (May 2022)"
-- each covering two distinct events with different dates, and two of them
appear in the new data. Keying on name alone would assign those articles to
the wrong disaster.

Rows are appended at the bottom rather than inserted in disaster order.
Inserting would shift ~1,700 existing rows and rewrite every row-relative
formula reference in the sheet; appending touches nothing that already
exists. Order does not matter downstream because All Disasters aggregates
with COUNTIF/AVERAGEIFS on the Disaster Row # column rather than by position.

Formula ranges are extended from row 1695 to the new last row. Every
Master-referencing formula in the workbook hardcodes 1695 as the end of the
data (2,817 references in All Disasters, 2 each in Active and Trash). Left
alone, the appended rows would be invisible to "# Articles Found",
"Avg Attribution Index" and "Contested Share" -- i.e. the disaster-level
Attribution Index would silently ignore every new article.

Columns left deliberately blank
-------------------------------
Relevance, Attribution Score/Category, D-Subtype, Contested and LLM
Rationale are blank because llm_coder.py selects exactly the rows whose
Relevance is still empty. Full Article Text, Archive URL, Archive Status
and Word Count are blank because archive.py skips rows that already carry
a status. Leaving them empty is what makes those two scripts pick these
rows up.

Engines is set to "RightLeaning" rather than the event's original engine
list, so the supplementary run stays identifiable in the data -- these
articles were collected through a different search engine than the rest of
the corpus, which is a distinction the methodology write-up depends on.
"""

import re
import shutil
from collections import defaultdict
from copy import copy
from datetime import datetime

import openpyxl
import pandas as pd

REPO = "/Users/hariaksha/Documents/GitHub/climate-disaster"
WORKBOOK = f"{REPO}/attribution.xlsx"
CLEAN_CSV = f"{REPO}/data/rightlean_articles_clean.csv"

OLD_LAST_ROW = 1695          # extent every Master! formula range currently assumes
ENGINE_TAG = "RightLeaning"

# Same category -> numeric mapping used for the existing rows and for the
# paper's AllSides robustness check.
ALLSIDES_MAP = {"left": -5.0, "lean left": -2.5, "center": 0.0,
                "lean right": 2.5, "right": 5.0}


def map_allsides(text):
    if not text:
        return None
    t = str(text).strip().lower()
    for key in ("lean left", "lean right", "left", "right", "center"):
        if t.startswith(key):
            return ALLSIDES_MAP[key]
    return None


def num(x):
    """Write integral floats as ints so new cells match existing formatting."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    if isinstance(x, float) and x.is_integer():
        return int(x)
    return x


def main():
    backup = f"{REPO}/attribution_backup_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    shutil.copy2(WORKBOOK, backup)
    print(f"Backup: {backup.rsplit('/',1)[-1]}\n")

    new = pd.read_csv(CLEAN_CSV)
    wb = openpyxl.load_workbook(WORKBOOK)          # formulas preserved
    ws, ad = wb["Master"], wb["All Disasters"]

    assert ws.max_row == OLD_LAST_ROW, f"Master ends at {ws.max_row}, expected {OLD_LAST_ROW}"

    # --- (Name, Begin Date) -> Disaster Row # -----------------------------
    key2row = {}
    for r in range(2, ad.max_row + 1):
        nm, bg = ad.cell(row=r, column=2).value, ad.cell(row=r, column=4).value
        if nm is not None:
            key2row[(str(nm).strip(), int(bg))] = ad.cell(row=r, column=1).value
    assert len(key2row) == ad.max_row - 1, "(Name, Begin Date) is not unique"
    print(f"Event key map: {len(key2row)} unique (Name, Begin Date) pairs")

    # --- next Article # per disaster --------------------------------------
    next_art = defaultdict(int)
    for r in range(2, ws.max_row + 1):
        d, a = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        if isinstance(d, int) and isinstance(a, int):
            next_art[d] = max(next_art[d], a)

    # --- domain -> lean values, from the Domain Ratings sheet -------------
    dr = wb["Domain Ratings"]
    bias, allsides = {}, {}
    for r in range(2, dr.max_row + 1):
        d = dr.cell(row=r, column=1).value
        if not d:
            continue
        af = dr.cell(row=r, column=5).value
        if af:
            m = re.search(r"Bias ([+-]?[0-9.]+)", af)
            if m:
                bias[d] = float(m.group(1))
        v = map_allsides(dr.cell(row=r, column=4).value)
        if v is not None:
            allsides[d] = v
    print(f"Lean lookup: {len(bias)} domains with a Bias score\n")

    # --- append ------------------------------------------------------------
    new = new.sort_values(["Name", "Begin Date", "Article #"]).reset_index(drop=True)
    start = OLD_LAST_ROW + 1
    written = 0
    unmatched = []

    for _, rec in new.iterrows():
        key = (str(rec["Name"]).strip(), int(rec["Begin Date"]))
        drow = key2row.get(key)
        if drow is None:
            unmatched.append(key)
            continue

        r = start + written
        next_art[drow] += 1
        dom = rec["Domain"]
        b = bias.get(dom)

        vals = {
            1: drow, 2: next_art[drow], 3: rec["Name"], 4: rec["Disaster"],
            5: int(rec["Begin Date"]), 6: int(rec["End Date"]),
            7: num(rec["CPI-Adjusted Cost"]), 8: num(rec["Unadjusted Cost"]),
            9: num(rec["Deaths"]), 10: rec["States"], 11: ENGINE_TAG,
            12: rec["URL"], 13: rec["Title"],
            14: None if pd.isna(rec["Snippet"]) else rec["Snippet"],
            # 15-20 left blank for llm_coder.py
            21: f'=IF(SUMPRODUCT(--ISNUMBER(SEARCH(\'National Domains\'!$A$2:$A$48,L{r})))>0,"Yes","No")',
            22: f'=IF(COUNTIFS(Calibration!$A:$A,A{r},Calibration!$B:$B,B{r})>0,"Yes","No")',
            # 23-26 left blank for archive.py
            27: dom,
            28: b,
            29: allsides.get(dom),
            30: ("Right-leaning" if b is not None and b > 0
                 else "Left-leaning" if b is not None else "Unidentified"),
        }

        for c, v in vals.items():
            cell = ws.cell(row=r, column=c, value=v)
            tmpl = ws.cell(row=2, column=c)
            cell.font = copy(tmpl.font)
            cell.alignment = copy(tmpl.alignment)
            cell.border = copy(tmpl.border)
            cell.fill = copy(tmpl.fill)
            cell.number_format = tmpl.number_format
        written += 1

    new_last = OLD_LAST_ROW + written
    print(f"Appended {written} rows -> Master now ends at {new_last}")
    if unmatched:
        print(f"  !! {len(unmatched)} rows had no matching event: {unmatched[:5]}")

    # --- extend every hardcoded Master! range ------------------------------
    pat = re.compile(r"(Master!\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?)" + str(OLD_LAST_ROW) + r"\b")
    patched = 0
    for sheet in wb.sheetnames:
        sh = wb[sheet]
        for row in sh.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "Master!" in c.value:
                    newv, n = pat.subn(r"\g<1>" + str(new_last), c.value)
                    if n:
                        c.value = newv
                        patched += n
    print(f"Extended {patched} Master! range references: {OLD_LAST_ROW} -> {new_last}")

    wb.save(WORKBOOK)
    print("Saved.")


if __name__ == "__main__":
    main()
