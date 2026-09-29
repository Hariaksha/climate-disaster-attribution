"""
Check retrieved article text in Master for the failure modes archive.py can
silently produce.

Every bug found so far had the same shape: the script reported success and
wrote something that was not an article. None of them raised an error, and all
were caught only by a person reading scrolling output. They are all cheap to
detect from the saved data:

  - CNN legacy transcript URLs saved the show's chronological index page
    (~8,000 words of unrelated segment titles) instead of a segment transcript
  - NPR audio-only pages saved audio-player markup ("Toggle more options",
    "<iframe src=", "Download Embed")
  - NPR /transcripts/{id} saved a two-sentence episode abstract (45-80 words)
    as though it were the full piece

so this reports on the data rather than the log.

Read-only: it never opens the workbook for writing and is safe to run while
archive.py is going.

    python3 code/validate_archive_output.py              # whole sheet
    python3 code/validate_archive_output.py --since 1696 # appended block only
    python3 code/validate_archive_output.py --verbose    # include text samples

Exit code is 1 when anything is flagged, so it can gate a run.
"""

import argparse
import re
import statistics
from collections import Counter, defaultdict
from urllib.parse import urlparse

import openpyxl

WORKBOOK = "/Users/hariaksha/Documents/GitHub/climate-disaster/attribution.xlsx"

# A real article in this corpus runs ~750 words (corpus median). Anything this
# short that was still accepted is a teaser, abstract or stub.
SHORT_TEXT_WORDS = 200

# Phrases that only appear in page furniture, never in article prose.
JUNK_MARKERS = [
    # "Return to Transcripts main page" is deliberately NOT here: it is a
    # breadcrumb on every CNN transcript page, genuine segment transcripts
    # included, so matching on it flags real content. Only the index page
    # carries the "continually updated" notice.
    ("This page is continually updated", "CNN chronological index page"),
    ("Toggle more options", "audio-player chrome"),
    ("NPR embedded audio player", "audio-player chrome"),
    ("<iframe src=", "raw embed markup"),
    ("Download Embed", "audio-player chrome"),
    ("Click here for a printable", "print-view furniture"),
]

# A domain whose stored text is far shorter than its own norm is usually a
# paywall serving previews for part of a run.
DOMAIN_OUTLIER_RATIO = 0.25
DOMAIN_MIN_SAMPLE = 8


def domain_of(url):
    return urlparse(str(url or "")).netloc.lower().removeprefix("www.")


def load_rows(since):
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb["Master"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    col = {name: header.index(name) for name in
           ("URL", "Title", "Relevance", "Full Article Text",
            "Archive Status", "Word Count")}
    rows = []
    r = 1
    for values in ws.iter_rows(min_row=2, values_only=True):
        r += 1
        if r < since:
            continue
        text = values[col["Full Article Text"]]
        if text is None or not str(text).strip():
            continue
        text = str(text)
        rows.append({
            "row": r,
            "url": values[col["URL"]],
            "title": values[col["Title"]],
            "relevance": values[col["Relevance"]],
            "status": values[col["Archive Status"]],
            "words": values[col["Word Count"]] or len(text.split()),
            "text": text,
            "domain": domain_of(values[col["URL"]]),
        })
    wb.close()
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--since", type=int, default=2,
                    help="first Master row to check (default: all)")
    ap.add_argument("--verbose", action="store_true",
                    help="show a text sample for each finding")
    args = ap.parse_args()

    rows = load_rows(args.since)
    if not rows:
        print("No rows with text in range.")
        return 0

    print(f"Checking {len(rows)} rows with text "
          f"(from row {args.since}) in Master\n")

    findings = defaultdict(list)

    # 1. page furniture that no article would contain
    for rec in rows:
        head = rec["text"][:4000]
        for marker, label in JUNK_MARKERS:
            if marker.lower() in head.lower():
                findings[f"Junk content: {label}"].append(rec)
                break

    # 2. accepted but too short to be an article
    for rec in rows:
        if rec["words"] < SHORT_TEXT_WORDS:
            findings[f"Short text (<{SHORT_TEXT_WORDS} words)"].append(rec)

    # 3. text saved on a row that is not marked found
    for rec in rows:
        if rec["status"] != "found":
            findings[f"Has text but status is {rec['status']!r}"].append(rec)

    # 4. per-domain length outliers (paywall previews mid-run)
    by_domain = defaultdict(list)
    for rec in rows:
        by_domain[rec["domain"]].append(rec)
    for dom, recs in by_domain.items():
        if len(recs) < DOMAIN_MIN_SAMPLE:
            continue
        median = statistics.median(r["words"] for r in recs)
        if median <= 0:
            continue
        for rec in recs:
            if rec["words"] < median * DOMAIN_OUTLIER_RATIO:
                findings[f"Far below {dom} median ({int(median)}w)"].append(rec)

    # 5. duplicate bodies -- the same wrong page saved for several rows
    by_text = defaultdict(list)
    for rec in rows:
        by_text[rec["text"][:600]].append(rec)
    for _, recs in by_text.items():
        if len(recs) > 1:
            findings["Identical text on multiple rows"].extend(recs)

    if not findings:
        print("No problems found.")
        return 0

    flagged = {r["row"] for recs in findings.values() for r in recs}
    for label, recs in sorted(findings.items(), key=lambda kv: -len(kv[1])):
        print(f"{label}  --  {len(recs)} row(s)")
        for rec in sorted(recs, key=lambda r: r["row"])[:12]:
            print(f"    row {rec['row']:5d}  {rec['words']:6d}w  "
                  f"{rec['domain']:24s} {str(rec['title'])[:44]}")
            if args.verbose:
                sample = " ".join(rec["text"].split())[:220]
                print(f"             {sample}")
        if len(recs) > 12:
            print(f"    ... and {len(recs) - 12} more")
        print()

    print(f"{len(flagged)} distinct row(s) flagged out of {len(rows)} checked.")
    print("Blank Full Article Text / Archive Status / Word Count on any row "
          "that is wrong; archive.py will re-queue it automatically.")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
